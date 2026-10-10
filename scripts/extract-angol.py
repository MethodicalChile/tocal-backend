#!/usr/bin/env python3
"""Extrae los datos de la obra EDS Aramco Angol desde las planillas Excel de Tocal a seed-assets/angol.json.
El comando seed_demo carga el JSON resultante."""
import json, re, sys, warnings, datetime as dt
from pathlib import Path
warnings.filterwarnings("ignore")
import openpyxl

D = Path(sys.argv[1] if len(sys.argv) > 1 else "../../Documentos/Informes_Obra")
OUT = Path(__file__).resolve().parent.parent / "seed-assets" / "angol.json"

def ws(file, sheet):
    return openpyxl.load_workbook(D / file, data_only=True)[sheet]
def iso(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, (dt.date, dt.datetime)) else None
def s(v):
    return None if v is None else (iso(v) if isinstance(v, (dt.date, dt.datetime)) else str(v).strip() or None)
def n(v):
    return float(v) if isinstance(v, (int, float)) else None

out = {}

w = ws("Lista de Chequeo PDR EDS Angol.xlsx", "TOCAL")
cat_map = {"1": "antecedentes_generales", "2": "antecedentes_legales", "3": "sso_medio_ambiente", "4": "terreno"}
items, sub, orden = [], None, 0
for r in range(44, w.max_row + 1):
    a, b, g, h = s(w.cell(r, 1).value), s(w.cell(r, 2).value), s(w.cell(r, 7).value), s(w.cell(r, 8).value)
    if a and re.match(r"^4\.\d+ [A-ZÁÉÍÓÚ]", a): sub = a
    elif a and re.match(r"^\d\. ", a): sub = None
    elif a and re.match(r"^\d+(\.\d+)+$", a) and b:
        orden += 1
        estado = {"cumple": "cumple", "no cumple": "no_cumple", "no aplica": "no_aplica"}.get((g or "cumple").lower(), "cumple")
        items.append({"codigo": a, "categoria": cat_map[a[0]], "subseccion": sub, "descripcion": b, "orden": orden, "estado": estado, "observacion": h})
out["checklist"] = {"dotacionDeclarada": n(w["L9"].value), "dotacionFemenina": n(w["L10"].value), "items": items}

w = ws("Estadisticas PDR Tocal- EDS Angol.xlsx", "PDR EDS ANGOL")
def rango(label, prev_end):
    m = re.match(r"(\d\d)-(\d\d)\s+al\s+(\d\d)-(\d\d)", label)
    d1, m1, d2, m2 = map(int, m.groups())
    y1 = 2025 if m1 == 12 else 2026
    ini = dt.date(y1, m1, d1)
    if prev_end and ini <= prev_end: ini = prev_end + dt.timedelta(days=1)  # etiqueta mal tipeada
    try: fin = dt.date(2025 if m2 == 12 else 2026, m2, d2)
    except ValueError: fin = ini + dt.timedelta(days=6)
    if fin < ini or (fin - ini).days > 6: fin = ini + dt.timedelta(days=6)
    return ini, fin
LABEL_FIX = {"01-06 al 07-05": "01-06 al 07-06", "29-06 al 05-08": "06-07 al 09-07"}  # etiquetas mal tipeadas en la planilla
semanas, prev = [], None
for r in range(8, w.max_row + 1):
    label = s(w.cell(r, 1).value)
    if not label or not re.match(r"\d\d-\d\d", label): continue
    label = LABEL_FIX.get(label, label)
    ini, fin = rango(label, prev); prev = fin
    v = [n(w.cell(r, c).value) or 0 for c in range(2, 19)]
    semanas.append({"numero": len(semanas) + 1, "etiqueta": label, "fechaInicio": ini.isoformat(), "fechaFin": fin.isoformat(),
        "personalOperativo": int(v[0]), "personalAdm": int(v[1]), "permisosLicencias": int(v[3]), "ingresos": int(v[4]), "finiquitos": int(v[5]),
        "hhOperativas": v[6], "hhAdm": v[7], "subcontratos": int(v[9]), "personalSc": int(v[10]), "hhSc": v[11],
        "incidentes": int(v[14]), "accidentes": int(v[15]), "diasPerdidos": int(v[16])})
PDF_EXTRA = [  # etiqueta, op, adm, perm, ing, fin, hhOp, hhAdm, subc, persSC, hhSC, inc, acc, dias
    ("10-07 al 16-07", 15,4,0,0,0, 540,144, 1,3,168, 0,0,0),
    ("17-07 al 23-07", 14,4,0,1,1, 588,168, 2,5,306, 0,0,0),
    ("24-07 al 30-07", 14,4,0,0,0, 588,168, 2,5,306, 0,0,0),
    ("31-07 al 06-08", 14,4,0,0,0, 546,168, 2,6,448, 0,0,0),
    ("07-08 al 13-08", 13,4,0,0,0, 546,168, 2,5,406, 0,0,0),
    ("14-08 al 20-08", 13,4,0,1,1, 546,168, 2,5,406, 0,0,0),
    ("21-08 al 27-08", 13,4,0,0,0, 546,168, 2,5,406, 0,0,0),
    ("28-08 al 03-09", 13,4,0,0,0, 546,168, 2,5,406, 0,0,0),
    ("04-09 al 10-09", 13,4,0,0,0, 546,168, 2,5,406, 0,1,0),
]
for (label, op, adm, perm, ing, fin_, hho, hha, subc, psc, hhsc, inc, acc, dias) in PDF_EXTRA:
    if any(x["etiqueta"] == label for x in semanas): continue
    ini, fin = rango(label, prev); prev = fin
    semanas.append({"numero": len(semanas) + 1, "etiqueta": label, "fechaInicio": ini.isoformat(), "fechaFin": fin.isoformat(),
        "personalOperativo": op, "personalAdm": adm, "permisosLicencias": perm, "ingresos": ing, "finiquitos": fin_,
        "hhOperativas": hho, "hhAdm": hha, "subcontratos": subc, "personalSc": psc, "hhSc": hhsc, "incidentes": inc, "accidentes": acc, "diasPerdidos": dias})
# Días perdidos: la columna semanal sólo registra los días del evento; la hoja ESTADISTICAS (mensual) y el
# informe (115 acumulados al 11-09) llevan los días de licencia continuados. Se distribuyen los totales mensuales.
META_DIAS = {"2026-04": 22, "2026-05": 34, "2026-06": 30, "2026-07": 9, "2026-08": 20}
for mes, meta in META_DIAS.items():
    ws_mes = [w for w in semanas if w["fechaInicio"].startswith(mes)]
    falta = meta - sum(w["diasPerdidos"] for w in ws_mes)
    for i, w in enumerate(ws_mes):
        if falta <= 0: break
        extra = falta // (len(ws_mes) - i)
        w["diasPerdidos"] += extra; falta -= extra
out["semanas"] = semanas

F = "OBRA EDS ARAMCO ANGOL - Informe semanal  - Semana del 07 al 11 de septiembre.xlsx"
def curva(sheet):
    w = ws(F, sheet); pts = []
    for r in range(7, 60):
        sem, f = w.cell(r, 2).value, w.cell(r, 3).value
        if not re.match(r"^\d+$", str(sem or "")) or not iso(f): continue
        pts.append({"semana": int(sem), "fecha": iso(f), "programado": n(w.cell(r, 5).value), "real": n(w.cell(r, 6).value), "reprogramado": n(w.cell(r, 7).value)})
    return pts
out["curvas"] = [
    {"variante": "general", "nombre": "Avance de obra — programa aprobado 27-01", "descripcion": "Control de avance se realiza con programación aprobada el 27-01", "orden": 1, "puntos": curva("Curva S Gral Aprobada")},
    {"variante": "accesos", "nombre": "Avance de obra con accesos (inicio julio)", "descripcion": "Control de avance con programa de accesos con inicio el 01-07", "orden": 3, "puntos": curva("Curva S Gral Aprobada Accesos")},
]
# Ruta crítica: GRAF. SEMANAL (fila 2 semana, 3 fecha, 4 programado, 5 real)
w = ws("260911_Avance Fisico Ruta Critica Aramco Angol.xlsx", "GRAF. SEMANAL")
pts = []
for c in range(2, w.max_column + 1):
    sem, f = w.cell(2, c).value, w.cell(3, c).value
    if not isinstance(sem, (int, float)) or not iso(f): continue
    pts.append({"semana": int(sem) + 1, "fecha": iso(f), "programado": n(w.cell(4, c).value), "real": n(w.cell(5, c).value), "reprogramado": None})
out["curvas"].insert(1, {"variante": "ruta_critica", "nombre": "Curva S ruta crítica", "descripcion": "Control avance ruta crítica", "orden": 2, "puntos": pts})

w = ws("260911_Avance Fisico Ruta Critica Aramco Angol.xlsx", "HOJA CONTROL ")
acts = []
for r in range(7, 60):
    raw = w.cell(r, 2).value
    if not isinstance(raw, str) or not raw.strip(): continue
    nivel = (len(raw) - len(raw.lstrip(" "))) // 3
    peso, pp, pa = n(w.cell(r, 9).value), n(w.cell(r, 6).value), n(w.cell(r, 7).value)
    st = s(w.cell(r, 10).value)
    acts.append({"orden": len(acts) + 1, "nivel": nivel, "descripcion": raw.strip(), "esResumen": peso is None and pp is None,
        "duracion": int(n(w.cell(r, 3).value) or 0) or None, "inicio": iso(w.cell(r, 4).value), "fin": iso(w.cell(r, 5).value),
        "peso": peso, "pctPrograma": pp, "pctActual": pa,
        "status": {"OK": "ok", "ATRASO": "atraso", "ADELANTO": "adelanto"}.get((st or "").upper()), "razonAccion": s(w.cell(r, 11).value)})
out["actividades"] = acts
out["rutaCriticaResumen"] = {"fechaChequeo": iso(w["D2"].value), "programa": n(w["G2"].value), "actual": n(w["G3"].value)}

w = ws(F, "EEPP")
out["obra"] = {"montoContractualUf": n(w["C9"].value), "anticipoUf": n(w["C10"].value)}
edp = []
for r in range(11, 30):
    lab = s(w.cell(r, 1).value)
    if not lab or not lab.upper().startswith("EEPP"): continue
    edp.append({"numero": int(re.sub(r"\D", "", lab)), "montoUf": n(w.cell(r, 3).value), "numeroFactura": s(w.cell(r, 4).value), "fecha": iso(w.cell(r, 5).value),
        "estado": "aprobado" if (s(w.cell(r, 6).value) or "").upper().startswith("APROB") else "presentado", "observacion": None})
out["estadosPago"] = edp
w = ws(F, "Adicionales")
ad = []
for r in range(9, 41):
    num, desc = w.cell(r, 1).value, s(w.cell(r, 2).value)
    if not isinstance(num, (int, float)) or not desc: continue
    est = (s(w.cell(r, 4).value) or "").lower()
    ad.append({"numero": int(num), "descripcion": desc, "fechaPresentacion": iso(w.cell(r, 3).value),
        "estado": "aprobado" if est.startswith("aprob") else "presentado" if est else "borrador",
        "clasificacion": s(w.cell(r, 5).value), "montoUf": n(w.cell(r, 6).value), "numeroOc": s(w.cell(r, 11).value), "numeroFactura": s(w.cell(r, 12).value), "observaciones": s(w.cell(r, 10).value)})
out["adicionales"] = ad

w = ws(F, "Consolidado 12.3.26")
rdi = []
for r in range(8, 40):
    num = w.cell(r, 1).value
    if not isinstance(num, (int, float)): continue
    p, rr, i = [bool(s(w.cell(r, c).value)) for c in (5, 6, 7)]
    estado = "respondida" if rr and not i else "incompleta" if i else "pendiente"
    fe, fr = iso(w.cell(r, 3).value), iso(w.cell(r, 8).value)
    if fe and fe.endswith("-12-04") and fe.startswith("2026"): fe = "2025" + fe[4:]   # año mal tipeado en la planilla
    if fr and fe and (fr < fe or fr > "2026-09-11"): fr = fe   # respuestas con año mal tipeado
    medio = s(w.cell(r, 2).value) or ""
    m = re.search(r"(RDI|ACL)\.?\s*n?°?\s*(\d+)", medio, re.I)
    rdi.append({"numero": int(num), "medio": re.sub(r"\s*(RDI|ACL)\.?\s*n?°?\s*\d+", "", medio, flags=re.I).strip(" .") or medio,
        "referencia": f"{m.group(1).upper()} N°{int(m.group(2)):02d}" if m else None,
        "fechaEnvio": fe, "materia": s(w.cell(r, 4).value), "estado": estado, "fechaRespuesta": fr if estado != "pendiente" else None,
        "especialidad": s(w.cell(r, 9).value), "respondidaPor": s(w.cell(r, 10).value), "observaciones": s(w.cell(r, 11).value)})
out["rdi"] = rdi

OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1))
print(f"{OUT}: checklist {len(items)} ítems, {len(semanas)} semanas, curvas {[len(c['puntos']) for c in out['curvas']]}, {len(acts)} actividades, {len(edp)} EDP, {len(ad)} adicionales, {len(rdi)} RDI")
