#!/usr/bin/env python3
"""Extrae el libro de ventas (hoja BD de 'DASHBOARD VENTAS TOCAL 2026 REV01.xlsx') a seed-assets/ventas.json."""
import json, sys, warnings, datetime as dt
from pathlib import Path
warnings.filterwarnings("ignore")
import openpyxl

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "../../Documentos/DASHBOARD VENTAS TOCAL 2026 REV01.xlsx")
OUT = Path(__file__).resolve().parent.parent / "seed-assets" / "ventas.json"

RUBRO = {"ENERGIA": "energia", "INDUSTRIAL": "industrial", "COMERCIAL": "comercial", "EDUCACIONAL": "educacional", "URBANIZACION": "urbanizacion", "OTROS": "otros"}
TIPO = {"FACTURA": "factura", "FACTURA ELECTRONICA": "factura_electronica", "NC": "nota_credito", "EXCENTA": "exenta", "EXENTA": "exenta"}

ws = openpyxl.load_workbook(SRC, data_only=True)["BD"]
docs = []
for r in range(2, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 15)]
    if not isinstance(v[6], (dt.date, dt.datetime)) or v[7] is None: continue
    s = lambda x: (str(x).strip() or None) if x is not None else None
    docs.append({
        "rutCliente": s(v[0]), "razonSocial": s(v[1]), "rubro": RUBRO.get((s(v[2]) or "OTROS").upper(), "otros"),
        "obra": s(v[3]), "tipoDocumento": TIPO.get((s(v[4]) or "FACTURA").upper(), "factura"), "numero": int(v[5]),
        "fecha": v[6].strftime("%Y-%m-%d"), "afecto": float(v[7]), "iva": float(v[8] or 0), "total": float(v[9] or 0), "ufDia": float(v[10]),
    })
OUT.write_text(json.dumps(docs, ensure_ascii=False, indent=0))
uf = sum(d["afecto"] / d["ufDia"] for d in docs)
print(f"{OUT}: {len(docs)} documentos, {len({d['rutCliente'] for d in docs})} clientes, total UF afecto {uf:,.2f}")
