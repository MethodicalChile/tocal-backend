from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .forms import VentaForm
from .models import Venta


@transaction.atomic
def import_workbook(file):
    payload = file.read()
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValidationError("La planilla descomprimida supera 100 MB.")
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValidationError("El archivo no es una planilla XLSX válida.") from exc
    if "BD" not in workbook.sheetnames:
        workbook.close()
        raise ValidationError("La planilla debe contener la hoja BD.")
    types = {
        "FACTURA": "factura",
        "FACTURA ELECTRONICA": "factura_electronica",
        "NC": "nota_credito",
        "EXCENTA": "exenta",
        "EXENTA": "exenta",
    }
    sectors = {key.upper(): key for key, _ in Venta._meta.get_field("rubro").choices}
    created = updated = processed = 0
    errors = []
    try:
        for number, row in enumerate(
            workbook["BD"].iter_rows(min_row=2, max_col=11, values_only=True), 2
        ):
            if number > 50000:
                raise ValidationError("La planilla supera 50000 filas.")
            if row[6] is None and row[7] is None:
                continue
            processed += 1
            day = row[6].date() if isinstance(row[6], datetime) else row[6]
            if isinstance(day, date):
                day = day.isoformat()
            data = dict(
                rut_cliente=str(row[0] or "").strip().upper().replace(".", ""),
                razon_social=str(row[1] or "").strip(),
                rubro=sectors.get(str(row[2]).strip().upper(), "otros"),
                obra=str(row[3] or "").strip(),
                tipo_documento=types.get(str(row[4]).strip().upper(), "factura"),
                numero=row[5],
                fecha=day,
                afecto=row[7],
                uf_dia=row[10],
            )
            try:
                doc_number = int(data["numero"])
            except (TypeError, ValueError):
                doc_number = None
            candidates = Venta.objects.filter(
                tipo_documento=data["tipo_documento"],
                numero=doc_number,
                rut_cliente=data["rut_cliente"],
            )
            if candidates.count() > 1:
                # La fecha distingue documentos históricos con el mismo número.
                try:
                    candidates = candidates.filter(fecha=date.fromisoformat(str(day)))
                except ValueError:
                    pass
            if candidates.count() > 1:
                errors.append(
                    f"Fila {number}: documento ambiguo, hay varios registros coincidentes."
                )
                continue
            existing = candidates.first()
            form = VentaForm(data, instance=existing, fetch_uf=False)
            if not form.is_valid():
                errors.append(
                    f"Fila {number}: "
                    + "; ".join(
                        f"{form.fields[k].label if k in form.fields else 'General'}: {' '.join(v)}"
                        for k, v in form.errors.items()
                    )
                )
            else:
                form.save()
                if existing:
                    updated += 1
                else:
                    created += 1
            if len(errors) >= 20:
                break
    finally:
        workbook.close()
    if errors:
        raise ValidationError(
            ["No se guardó la importación. Corrige los errores y vuelve a intentar.", *errors]
        )
    if not processed:
        raise ValidationError("La hoja BD no contiene documentos.")
    return created, updated
